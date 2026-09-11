"""V/sigma study explainer - plain language, python-docx."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\hoche\OneDrive\Desktop\Cluster Analysis HTML\vsigma_study\VSigma_Study_Explained.docx"
RANK = r"C:\Users\hoche\OneDrive\Desktop\Cluster Analysis HTML\vsigma_study\vsigma_ranking.png"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

doc.add_heading("How We Measured Cluster Rotation (V/σ)", 0)
p = doc.add_paragraph()
r = p.add_run("A plain-language guide to our study of 24 globular clusters, using data from the Gaia space telescope.")
r.italic = True

doc.add_heading("1. The question", level=1)
doc.add_paragraph(
    "A globular cluster is a ball of hundreds of thousands of stars. Do those stars just "
    "buzz around randomly, or does the whole ball also spin? We measure this with one "
    "number: V/σ (“V over sigma”). V is how fast the cluster spins in an organized way; "
    "σ (sigma) is how fast stars shuffle around randomly. V/σ near 0 means no spin; "
    "higher values (0.3–0.6) mean the cluster clearly rotates.")

doc.add_heading("2. Getting the stars", level=1)
for t in [
    "We downloaded every star Gaia sees toward each cluster (out to the cluster’s edge, its “tidal radius”).",
    "Cluster members all drift across the sky together (same proper motion), crowd toward the cluster "
    "center, and line up on one narrow track in a color-vs-brightness chart (because they were born "
    "together). Field stars do none of these.",
    "Each star got a membership probability from these three tests. For this study we kept only stars "
    "above 90% probability.",
]:
    doc.add_paragraph(t, style="List Bullet")

doc.add_heading("3. Measuring the spin (V)", level=1)
doc.add_paragraph(
    "For each star we take its sideways drift relative to the cluster’s average motion, and split it "
    "into “toward/away from the center” and “around the center”. If the cluster spins, stars all "
    "around one ring drift the same way around the center — so we average that “around” motion in "
    "rings of increasing radius. The largest ring average is V.")
doc.add_paragraph(
    "One trap we avoided: if you slightly mismeasure the cluster’s average motion, you get a fake "
    "pattern that looks like rotation but flips sign around the ring (a sinusoid). Real rotation is the "
    "same all the way around. We tested for the fake pattern separately — it also turned out to be a "
    "great detector of contaminated samples.")

doc.add_heading("4. Measuring the shuffle (σ)", level=1)
doc.add_paragraph(
    "σ is the spread of star speeds — but Gaia’s measurement errors add fake spread. We subtracted "
    "the known measurement error to recover the true spread. Skipping this step makes σ too big and "
    "V/σ misleadingly tiny — for distant clusters this correction changes the answer completely.")

doc.add_heading("5. Honesty checks", level=1)
for t in [
    "Detection test: a statistical test (chi-squared) asks “could these ring averages be pure noise?” "
    "Only clusters passing at 99% confidence count as rotating.",
    "Quality flags: every cluster is graded good / error-dominated / contaminated / unmeasurable. Two "
    "distant clusters (NGC 2419, Pal 5) are honestly “unmeasurable” — Gaia’s errors exceed their "
    "entire internal motion.",
    "Validation: clusters already known to rotate (from published papers) all ranked at the top of our "
    "list, and the known non-rotator NGC 288 failed the detection test — exactly as it should.",
]:
    doc.add_paragraph(t, style="List Bullet")

doc.add_heading("6. What we found", level=1)
for t in [
    "Fastest relative rotators: 47 Tuc (V/σ = 0.63) and ω Centauri (0.52) — matching published values.",
    "Bigger clusters spin faster relative to their shuffle: V/σ rises with cluster size (correlation "
    "0.71, statistically significant). Likely because small dense clusters “forget” their spin faster "
    "through star-to-star encounters.",
    "Metal content ([Fe/H]) showed no correlation with V/σ in our sample.",
    "NGC 6266 (M62) shows a strong rotation signal that past surveys never tested — our most original "
    "result (with a caution: its outskirts are contaminated by bulge stars).",
]:
    doc.add_paragraph(t, style="List Bullet")

doc.add_picture(RANK, width=Inches(4.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
r = cap.add_run("All 24 clusters ranked by V/σ. Blue = confirmed rotators in past papers; they rank high in our measurement too.")
r.italic = True
r.font.size = Pt(9)
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("7. Where everything lives", level=1)
doc.add_paragraph(
    "Folder “vsigma_study” inside Cluster Analysis HTML: the full technical write-up (METHODS.md), "
    "the analysis program (vsigma_pipeline.py), the results table (vsigma_summary.csv), and per-cluster "
    "plots. The interactive version is in the app — open the V/σ study screen from the top bar.")

doc.save(OUT)
print("wrote", OUT)
